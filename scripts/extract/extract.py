"""Turn input documents into plain text the agent can read (FR-INT-01, FR-INP-02).

Supported: .md, .txt, .docx, .pdf, .pptx. Anything else is listed as skipped.

Usage:
    python scripts/extract/extract.py <file>            # print the text of one file
    python scripts/extract/extract.py --launch <slug>   # extract every input of a launch

For a launch, text goes to launches/<slug>/intake/extracted/ (one .txt per input)
and a summary to launches/<slug>/intake/manifest.yaml. These live outside inputs/
so extracting never changes the inputs themselves.
"""

from __future__ import annotations

import sys
from dataclasses import dataclass
from pathlib import Path

if __package__ in (None, ""):
    sys.path.insert(0, str(Path(__file__).resolve().parents[2]))

from scripts.util import LAUNCHES_DIR, dump_yaml, now_iso

SUPPORTED = {".md", ".markdown", ".txt", ".docx", ".pdf", ".pptx"}
IGNORED_NAMES = {".DS_Store", ".gitkeep", "Thumbs.db"}
# Below this many characters per page, a PDF is probably a scanned image.
MIN_PDF_CHARS_PER_PAGE = 40
W_NS = "http://schemas.openxmlformats.org/wordprocessingml/2006/main"


@dataclass
class Extracted:
    path: Path
    text: str
    status: str  # "ok" | "warning" | "skipped" | "error"
    note: str = ""


def _docx_text(path: Path) -> str:
    from docx import Document

    doc = Document(str(path))
    parts: list[str] = []
    body = doc.element.body
    # Walk the body in order so tables stay where they appear in the document.
    for child in body.iterchildren():
        tag = child.tag.rsplit("}", 1)[-1]
        if tag == "p":
            text = "".join(t.text or "" for t in child.iter() if t.tag.endswith("}t"))
            style = child.find(".//{*}pStyle")
            style_name = style.get(f"{{{W_NS}}}val") if style is not None else ""
            if text.strip():
                if style_name and style_name.lower().startswith("heading"):
                    level = "".join(ch for ch in style_name if ch.isdigit()) or "1"
                    parts.append("#" * min(int(level), 6) + " " + text.strip())
                else:
                    parts.append(text)
        elif tag == "tbl":
            for row in child.iter("{*}tr"):
                cells = []
                for cell in row.iter("{*}tc"):
                    paragraphs = ("".join(t.text or "" for t in p.iter("{*}t")) for p in cell.iter("{*}p"))
                    cells.append(" ".join(paragraphs).strip())
                parts.append("| " + " | ".join(cells) + " |")
    return "\n".join(parts)


def _pdf_text(path: Path) -> tuple[str, str]:
    from pypdf import PdfReader

    reader = PdfReader(str(path))
    page_texts = [(page.extract_text() or "").strip() for page in reader.pages]
    text = "\n\n".join(f"[page {i}]\n{t}" for i, t in enumerate(page_texts, start=1))
    note = ""
    if page_texts and sum(map(len, page_texts)) < MIN_PDF_CHARS_PER_PAGE * len(page_texts):
        note = "very little text found: this PDF may be a scanned image. Ask the PMM for a text version."
    return text, note


def _pptx_text(path: Path) -> str:
    from pptx import Presentation

    prs = Presentation(str(path))
    parts = []
    for i, slide in enumerate(prs.slides, start=1):
        lines = []
        for shape in slide.shapes:
            if shape.has_text_frame:
                text = shape.text_frame.text.strip()
                if text:
                    lines.append(text)
            if getattr(shape, "has_table", False) and shape.has_table:
                for row in shape.table.rows:
                    lines.append("| " + " | ".join(c.text.strip() for c in row.cells) + " |")
        if slide.has_notes_slide:
            notes = slide.notes_slide.notes_text_frame.text.strip()
            if notes:
                lines.append(f"[speaker notes] {notes}")
        parts.append(f"[slide {i}]\n" + "\n".join(lines))
    return "\n\n".join(parts)


def extract_text(path: Path) -> Extracted:
    suffix = path.suffix.lower()
    if suffix not in SUPPORTED:
        return Extracted(path, "", "skipped", f"{suffix or 'this'} files are not read")
    try:
        if suffix in (".md", ".markdown", ".txt"):
            return Extracted(path, path.read_text(encoding="utf-8", errors="replace"), "ok")
        if suffix == ".docx":
            return Extracted(path, _docx_text(path), "ok")
        if suffix == ".pdf":
            text, note = _pdf_text(path)
            return Extracted(path, text, "warning" if note else "ok", note)
        return Extracted(path, _pptx_text(path), "ok")
    except Exception as exc:  # a broken file should never stop intake
        return Extracted(path, "", "error", f"could not be read ({type(exc).__name__}: {exc})")


def input_files(launch_dir: Path) -> list[Path]:
    inputs = launch_dir / "inputs"
    return sorted(p for p in inputs.rglob("*") if p.is_file() and p.name not in IGNORED_NAMES)


def extract_launch(launch_dir: Path) -> list[Extracted]:
    out_dir = launch_dir / "intake" / "extracted"
    out_dir.mkdir(parents=True, exist_ok=True)
    for old in out_dir.glob("*.txt"):
        old.unlink()
    results = []
    manifest = []
    for path in input_files(launch_dir):
        rel = path.relative_to(launch_dir).as_posix()
        result = extract_text(path)
        results.append(result)
        entry = {"file": rel, "status": result.status, "characters": len(result.text)}
        if result.note:
            entry["note"] = result.note
        if result.text:
            target = out_dir / (rel.replace("/", "__") + ".txt")
            target.write_text(result.text, encoding="utf-8")
            entry["text_file"] = target.relative_to(launch_dir).as_posix()
        manifest.append(entry)
    dump_yaml({"extracted_at": now_iso(), "files": manifest}, launch_dir / "intake" / "manifest.yaml")
    return results


def main(argv: list[str]) -> int:
    if not argv:
        print(__doc__)
        return 2
    if argv[0] == "--launch":
        launch_dir = LAUNCHES_DIR / argv[1]
        if not (launch_dir / "inputs").is_dir():
            print(f"[X] No launch called '{argv[1]}' with an inputs/ folder.")
            return 1
        results = extract_launch(launch_dir)
        for r in results:
            rel = r.path.relative_to(launch_dir).as_posix()
            mark = {"ok": "[OK]", "warning": "[!]", "skipped": "[-]", "error": "[X]"}[r.status]
            print(f"{mark} {rel}" + (f": {r.note}" if r.note else ""))
        print(f"Text saved in launches/{argv[1]}/intake/extracted/")
        return 0
    result = extract_text(Path(argv[0]))
    if result.status in ("skipped", "error"):
        print(f"[X] {result.note}")
        return 1
    print(result.text)
    return 0


if __name__ == "__main__":
    sys.exit(main(sys.argv[1:]))
