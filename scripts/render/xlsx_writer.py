"""The messaging grid as an Excel file (FR-MSG-03, FR-RND-02).

One row per industry × persona, with filters on every column and the header row
frozen. Uses only features that survive import into Google Sheets: plain cells,
fills, wrap text, an autofilter and frozen panes. No macros, formulas or charts.
"""

from __future__ import annotations

from datetime import datetime, timezone
from pathlib import Path

from openpyxl import Workbook
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
from openpyxl.utils import get_column_letter

COLUMNS = [
    ("Industry", 22),
    ("Persona", 24),
    ("Buying role", 18),
    ("Headline", 42),
    ("Pain addressed", 42),
    ("Pillar", 8),
    ("Pillar headline", 30),
    ("Value hypothesis", 36),
    ("Label", 20),
    ("Proof", 30),
]


def _hex(color: str) -> str:
    return color.lstrip("#").upper()


def write_grid(path: Path, messaging: dict, *, industries: dict, personas: dict, pain_text, brand: dict,
               status_line: str, footer: str, render_date: str) -> Path:
    colors, fonts = brand["colors"], brand["fonts"]
    body_font = fonts.get("body", "Arial")
    pillars = {p["id"]: p["headline"] for p in messaging.get("pillars", [])}

    wb = Workbook()
    ws = wb.active
    ws.title = "Messaging grid"
    thin = Side(style="thin", color=_hex(colors.get("secondary") or colors["text"]))
    border = Border(left=thin, right=thin, top=thin, bottom=thin)
    header_font = Font(name=body_font, bold=True, color=_hex(colors["background"]))
    header_fill = PatternFill("solid", fgColor=_hex(colors["primary"]))
    cell_font = Font(name=body_font, color=_hex(colors["text"]))
    wrap = Alignment(wrap_text=True, vertical="top")

    for col, (name, width) in enumerate(COLUMNS, start=1):
        cell = ws.cell(row=1, column=col, value=name)
        cell.font, cell.fill, cell.alignment, cell.border = header_font, header_fill, wrap, border
        ws.column_dimensions[get_column_letter(col)].width = width

    for r, row in enumerate(messaging.get("grid", []), start=2):
        persona = personas.get(row["persona_id"], {})
        values = [
            industries.get(row["industry"], {}).get("name", row["industry"]),
            persona.get("role", row["persona_id"]),
            persona.get("buying_role", "").replace("-", " "),
            row["headline"],
            pain_text(row["pain_ref"]),
            row["pillar_id"],
            pillars.get(row["pillar_id"], ""),
            row["value_hypothesis"]["text"],
            row["value_hypothesis"]["label"],
            row["proof"],
        ]
        for col, value in enumerate(values, start=1):
            cell = ws.cell(row=r, column=col, value=value)
            cell.font, cell.alignment, cell.border = cell_font, wrap, border

    last_row = max(ws.max_row, 1)
    ws.auto_filter.ref = f"A1:{get_column_letter(len(COLUMNS))}{last_row}"
    ws.freeze_panes = "A2"
    ws.oddFooter.center.text = footer
    ws.oddHeader.right.text = status_line

    about = wb.create_sheet("About")
    for i, line in enumerate([
        "Messaging grid",
        status_line,
        footer,
        "Value hypotheses are labeled [ASSUMPTION] unless backed by customer data.",
        "Edit messaging in messaging.yaml (or ask the agent); edits to this spreadsheet are not copied back.",
    ], start=1):
        cell = about.cell(row=i, column=1, value=line)
        cell.font = Font(name=body_font, bold=(i == 1), color=_hex(colors["text"]))
    about.column_dimensions["A"].width = 110

    stamp = datetime.fromisoformat(render_date).replace(tzinfo=timezone.utc).replace(tzinfo=None)
    wb.properties.creator = "Industry Launch Accelerator"
    wb.properties.lastModifiedBy = "Industry Launch Accelerator"
    wb.properties.title = "Messaging grid"
    wb.properties.created = stamp
    wb.properties.modified = stamp
    path.parent.mkdir(parents=True, exist_ok=True)
    wb.save(str(path))
    return path
