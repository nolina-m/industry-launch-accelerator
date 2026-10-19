"""Write a Word document from the toolkit's Markdown, styled by brand.yaml (FR-RND-01, FR-OUT-02).

Supports the small Markdown subset the asset templates produce: headings (# to ####),
paragraphs, bullets (two levels), numbered lists, quotes (>), tables, images on
their own line, **bold** and *italic*. Only features that survive import into
Google Docs are used: built-in styles, simple tables, inline PNG images, plain
header and footer text. No SmartArt, macros or embedded fonts.
"""

from __future__ import annotations

import re
from datetime import datetime, timezone
from pathlib import Path

from docx import Document
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Cm, Pt, RGBColor

from scripts.brand.inspect_template import open_document

DEFAULT_STYLES = {
    "title": "Title",
    "heading1": "Heading 1",
    "heading2": "Heading 2",
    "heading3": "Heading 3",
    "heading4": "Heading 4",
    "body": "Normal",
    "list_bullet": "List Bullet",
    "list_bullet2": "List Bullet 2",
    "list_number": "List Number",
    "quote": "Quote",
    "table": "Table Grid",
}
INLINE_RE = re.compile(r"(\*\*[^*]+\*\*|\*[^*\s][^*]*\*|_[^_\s][^_]*_|`[^`]+`)")
IMAGE_RE = re.compile(r"^!\[([^\]]*)\]\(([^)]+)\)$")
TABLE_SEP_RE = re.compile(r"^\|?\s*:?-{3,}")
ALIGN = {"header-left": WD_ALIGN_PARAGRAPH.LEFT, "header-right": WD_ALIGN_PARAGRAPH.RIGHT,
         "header-center": WD_ALIGN_PARAGRAPH.CENTER}


def _rgb(hex_color: str) -> RGBColor:
    return RGBColor.from_string(hex_color.lstrip("#").upper())


def _split_row(line: str) -> list[str]:
    line = line.strip()
    if line.startswith("|"):
        line = line[1:]
    if line.endswith("|") and not line.endswith("\\|"):
        line = line[:-1]
    cells = re.split(r"(?<!\\)\|", line)
    return [c.strip().replace("\\|", "|") for c in cells]


class DocxWriter:
    def __init__(self, brand: dict, launch_dir: Path, header_label: str = "", footer: str = "",
                 render_date: str | None = None):
        self.brand = brand
        self.launch_dir = launch_dir
        self.header_label = header_label
        self.footer = footer
        self.render_date = render_date
        template = (brand.get("word_template") or {}).get("path")
        self.from_template = bool(template) and (launch_dir / template).is_file()
        if self.from_template:
            self.doc = open_document(launch_dir / template)
            self._clear_body()
        else:
            self.doc = Document()
            self._use_current_word_format()
            self._apply_theme()
        self.styles = dict(DEFAULT_STYLES)
        if self.from_template:
            self.styles.update(brand["word_template"].get("style_map", {}))
        self._setup_header_footer()

    # ------------------------------------------------------------------ setup
    def _clear_body(self) -> None:
        body = self.doc.element.body
        for child in list(body):
            if child.tag != qn("w:sectPr"):
                body.remove(child)

    def _use_current_word_format(self) -> None:
        """python-docx's base file says "Word 2010" (compatibility mode 14), so Word shows
        "[Compatibility Mode]". Mode 15 is the current format. A PMM's own template keeps its setting."""
        for el in self.doc.settings.element.iter(qn("w:compatSetting")):
            if el.get(qn("w:name")) == "compatibilityMode":
                el.set(qn("w:val"), "15")

    def _apply_theme(self) -> None:
        """Neutral or brand theme on Word's built-in styles (FR-INP-10, FR-RND-06)."""
        colors, fonts = self.brand["colors"], self.brand["fonts"]
        normal = self.doc.styles["Normal"]
        normal.font.name = fonts["body"]
        normal.font.size = Pt(10.5)
        normal.font.color.rgb = _rgb(colors["text"])
        self._set_east_asian_font(normal, fonts["body"])
        for name, size in (("Title", 24), ("Heading 1", 16), ("Heading 2", 13), ("Heading 3", 11.5),
                           ("Heading 4", 10.5)):
            style = self.doc.styles[name]
            style.font.name = fonts["heading"]
            style.font.size = Pt(size)
            style.font.color.rgb = _rgb(colors["primary"] if name != "Heading 4" else colors["text"])
            self._set_east_asian_font(style, fonts["heading"])
        for name in ("List Bullet", "List Bullet 2", "List Number", "Quote", "Table Grid"):
            style = self.doc.styles[name]
            style.font.name = fonts["body"]
            style.font.color.rgb = _rgb(colors["text"])
            self._set_east_asian_font(style, fonts["body"])
        for name in ("Normal", "Title", "Heading 1", "Heading 2", "Heading 3", "Heading 4", "List Bullet",
                     "List Bullet 2", "List Number", "Quote", "Table Grid"):
            self._drop_theme_colors(self.doc.styles[name], colors["primary"])
        normal.paragraph_format.space_after = Pt(5)
        normal.paragraph_format.line_spacing = 1.1
        for name in ("Heading 1", "Heading 2", "Heading 3", "Heading 4"):
            self.doc.styles[name].paragraph_format.space_before = Pt(12)
            self.doc.styles[name].paragraph_format.space_after = Pt(4)
        for name in ("List Bullet", "List Bullet 2", "List Number"):
            self.doc.styles[name].paragraph_format.space_after = Pt(2)
        for section in self.doc.sections:
            section.left_margin = section.right_margin = Cm(2.2)

    @staticmethod
    def _set_east_asian_font(style, font_name: str) -> None:
        """Set the font for every script, and remove theme fonts, which Word would use instead."""
        rpr = style.element.get_or_add_rPr()
        rfonts = rpr.find(qn("w:rFonts"))
        if rfonts is None:
            rfonts = OxmlElement("w:rFonts")
            rpr.append(rfonts)
        for attr in ("w:asciiTheme", "w:hAnsiTheme", "w:eastAsiaTheme", "w:cstheme"):
            rfonts.attrib.pop(qn(attr), None)
        for attr in ("w:ascii", "w:hAnsi", "w:cs", "w:eastAsia"):
            rfonts.set(qn(attr), font_name)

    @staticmethod
    def _drop_theme_colors(style, border_color: str) -> None:
        """Word applies theme colors before plain ones, so remove them from styles we brand.
        Border lines (like the rule under the title) take the brand's primary color."""
        for el in style.element.iter():
            for attr in ("w:themeColor", "w:themeShade", "w:themeTint"):
                el.attrib.pop(qn(attr), None)
            if el.tag in (qn("w:top"), qn("w:bottom"), qn("w:left"), qn("w:right")) and el.get(qn("w:color")) \
                    not in (None, "auto"):
                el.set(qn("w:color"), border_color.lstrip("#").upper())

    def _setup_header_footer(self) -> None:
        section = self.doc.sections[0]
        header = section.header
        header.is_linked_to_previous = False
        para = header.paragraphs[0] if header.paragraphs else header.add_paragraph()
        for run in list(para.runs):
            run.text = ""
        logo = (self.brand.get("logo") or {}).get("light_bg")
        placement = (self.brand.get("logo") or {}).get("placement", "none")
        if logo and placement != "none" and (self.launch_dir / logo).is_file():
            para.alignment = ALIGN.get(placement, WD_ALIGN_PARAGRAPH.RIGHT)
            para.add_run().add_picture(str(self.launch_dir / logo), height=Cm(1.0))
            if self.header_label:
                para = header.add_paragraph()
        if self.header_label:
            run = para.add_run(self.header_label)
            run.font.size = Pt(8)
            run.font.color.rgb = _rgb(self.brand["colors"]["secondary"] if "secondary" in self.brand["colors"]
                                      else self.brand["colors"]["text"])
        footer = section.footer
        footer.is_linked_to_previous = False
        fpara = footer.paragraphs[0] if footer.paragraphs else footer.add_paragraph()
        for run in list(fpara.runs):
            run.text = ""
        if self.footer:
            run = fpara.add_run(self.footer)
            run.font.size = Pt(8)
            run.font.color.rgb = _rgb(self.brand["colors"]["text"])

    # ------------------------------------------------------------------ writing
    def _style(self, key: str):
        name = self.styles.get(key) or DEFAULT_STYLES.get(key)
        try:
            return self.doc.styles[name]
        except KeyError:
            return self.doc.styles[self.styles.get("body", "Normal")] if key != "body" else None

    def _add_inline(self, para, text: str, bold: bool = False) -> None:
        for part in INLINE_RE.split(text):
            if not part:
                continue
            if part.startswith("**") and part.endswith("**"):
                run = para.add_run(part[2:-2])
                run.bold = True
            elif (part.startswith("*") and part.endswith("*")) or (part.startswith("_") and part.endswith("_")):
                run = para.add_run(part[1:-1])
                run.italic = True
            elif part.startswith("`") and part.endswith("`"):
                run = para.add_run(part[1:-1])
            else:
                run = para.add_run(part)
            if bold:
                run.bold = True

    def paragraph(self, text: str, style_key: str = "body") -> None:
        para = self.doc.add_paragraph(style=self._style(style_key))
        self._add_inline(para, text)

    def table(self, rows: list[list[str]]) -> None:
        width = max(len(r) for r in rows)
        rows = [r + [""] * (width - len(r)) for r in rows]
        table = self.doc.add_table(rows=len(rows), cols=width)
        style = self._style("table")
        if style is not None:
            table.style = style
        self._borders(table)
        header_fill = self.brand["colors"]["primary"].lstrip("#").upper()
        header_text = self.brand["colors"]["background"]
        for r, row in enumerate(rows):
            for c, value in enumerate(row):
                cell = table.cell(r, c)
                cell.text = ""
                para = cell.paragraphs[0]
                para.paragraph_format.space_after = Pt(2)
                self._add_inline(para, value, bold=(r == 0))
                if r == 0:
                    shading = OxmlElement("w:shd")
                    shading.set(qn("w:val"), "clear")
                    shading.set(qn("w:color"), "auto")
                    shading.set(qn("w:fill"), header_fill)
                    cell._tc.get_or_add_tcPr().append(shading)
                    for run in para.runs:
                        run.font.color.rgb = _rgb(header_text)
        self._column_widths(table, rows)
        self.doc.add_paragraph(style=self._style("body"))

    def _column_widths(self, table, rows: list[list[str]]) -> None:
        """Share the page width between columns by how much text each holds."""
        section = self.doc.sections[-1]
        usable = section.page_width - section.left_margin - section.right_margin
        # Bold header text is wider than body text, and short IDs like "WS-1" shouldn't wrap.
        weights = [min(max(max(len(r[c]) for r in rows[1:]) if len(rows) > 1 else 0, len(rows[0][c]) + 3, 8), 60)
                   for c in range(len(rows[0]))]
        total = sum(weights)
        widths = [int(usable * w / total) for w in weights]
        table.autofit = False
        grid = table._tbl.tblGrid
        for col, w in zip(grid.findall(qn("w:gridCol")), widths, strict=False):
            col.set(qn("w:w"), str(int(w / 635)))  # EMU to twentieths of a point
        for row in table.rows:
            for cell, w in zip(row.cells, widths, strict=False):
                cell.width = w

    def _borders(self, table) -> None:
        """Explicit thin borders, so tables keep their lines in Google Docs and other viewers."""
        color = self.brand["colors"].get("secondary") or self.brand["colors"]["text"]
        tbl_pr = table._tbl.tblPr
        borders = OxmlElement("w:tblBorders")
        for edge in ("top", "left", "bottom", "right", "insideH", "insideV"):
            el = OxmlElement(f"w:{edge}")
            el.set(qn("w:val"), "single")
            el.set(qn("w:sz"), "4")
            el.set(qn("w:space"), "0")
            el.set(qn("w:color"), color.lstrip("#").upper())
            borders.append(el)
        tbl_pr.append(borders)
        width = OxmlElement("w:tblW")  # full page width, so text wraps less
        width.set(qn("w:w"), "5000")
        width.set(qn("w:type"), "pct")
        for old in tbl_pr.findall(qn("w:tblW")):
            tbl_pr.remove(old)
        tbl_pr.append(width)

    def image(self, rel_path: str, alt: str) -> None:
        path = (self.launch_dir / "outputs" / rel_path) if not Path(rel_path).is_absolute() else Path(rel_path)
        if not path.is_file():
            self.paragraph(f"[Image missing: {alt}]")
            return
        self.doc.add_picture(str(path), width=Cm(15.5))
        last = self.doc.paragraphs[-1]
        last.alignment = WD_ALIGN_PARAGRAPH.CENTER

    def add_markdown(self, markdown: str) -> None:
        lines = markdown.splitlines()
        i = 0
        buffer: list[str] = []
        first_h1 = True

        def flush():
            if buffer:
                self.paragraph(" ".join(s.strip() for s in buffer))
                buffer.clear()

        while i < len(lines):
            line = lines[i]
            stripped = line.strip()
            if not stripped or stripped == "---":
                flush()
            elif stripped.startswith("#"):
                flush()
                level = len(stripped) - len(stripped.lstrip("#"))
                text = stripped[level:].strip()
                if level == 1 and first_h1:
                    self.paragraph(text, "title")
                    first_h1 = False
                else:
                    self.paragraph(text, f"heading{min(level, 4)}" if level > 1 else "heading1")
            elif stripped.startswith("|"):
                flush()
                rows = []
                while i < len(lines) and lines[i].strip().startswith("|"):
                    if not TABLE_SEP_RE.match(lines[i].strip()):
                        rows.append(_split_row(lines[i]))
                    i += 1
                self.table(rows)
                continue
            elif IMAGE_RE.match(stripped):
                flush()
                alt, path = IMAGE_RE.match(stripped).groups()
                self.image(path, alt)
            elif re.match(r"^(\s*)[-*] ", line):
                flush()
                indent = len(line) - len(line.lstrip())
                self.paragraph(re.sub(r"^\s*[-*] ", "", line), "list_bullet2" if indent >= 2 else "list_bullet")
            elif re.match(r"^\s*\d+[.)] ", line):
                flush()
                self.paragraph(re.sub(r"^\s*\d+[.)] ", "", line), "list_number")
            elif stripped.startswith(">"):
                flush()
                self.paragraph(stripped.lstrip("> ").strip(), "quote")
            else:
                buffer.append(stripped)
            i += 1
        flush()

    def save(self, path: Path, title: str) -> None:
        props = self.doc.core_properties
        props.title = title
        props.author = "Industry Launch Accelerator"
        props.last_modified_by = "Industry Launch Accelerator"
        props.comments = self.footer
        props.revision = 1
        stamp = datetime.fromisoformat(self.render_date) if self.render_date else datetime.now(timezone.utc)
        if stamp.tzinfo is None:
            stamp = stamp.replace(tzinfo=timezone.utc)
        props.created = props.modified = stamp
        path.parent.mkdir(parents=True, exist_ok=True)
        self.doc.save(str(path))
