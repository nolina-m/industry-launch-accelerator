"""Rendering YAML to Markdown and Word (FR-RND-01..06, FR-GRD-06, FR-INP-10, FR-INP-11, FR-STA-03)."""

import zipfile

import pytest
from docx import Document

from scripts.render.render import FOOTER, Renderer, cited_in, labels_in
from scripts.sources import add as add_source
from scripts.state.launch_state import LaunchState
from scripts.util import dump_yaml, load_yaml
from tests import factories as f
from tests.docbuilders import make_word_template
from tests.test_state import INDUSTRIES, write_outputs

ALL_BUT_QA = ["intake", "research:ind-a", "research:ind-b", "competitive", "strategy",
              "personas:ind-a", "personas:ind-b", "messaging", "playbook:ind-a", "playbook:ind-b",
              "demo:ind-a", "demo:ind-b", "web:ind-a", "web:ind-b"]


@pytest.fixture
def launch(tmp_path):
    s = LaunchState.create(tmp_path / "r", "r", "Render test")
    for i in INDUSTRIES:
        s.data["industries"].append({"id": i, "name": i.upper()})
    s.data["competitors"].append({"id": "comp-a", "name": "Competitor A", "suggested": False})
    s.refresh()
    for unit in ALL_BUT_QA:
        s.start(unit)
        write_outputs(s.dir, unit)
        s.submit(unit)
        s.approve(unit)
    s.save()
    add_source(s.dir, url="https://example.org/a", title="Source title", publisher="Publisher", tier=1,
               excerpt="supporting text", published="2026-01-01")
    return s


def render(launch):
    return Renderer(launch.dir).render_all()


def test_all_eight_assets_render_for_two_industries(launch):
    """M4 acceptance: all 8 assets render for 2+ industries."""
    results = render(launch)
    names = sorted(r.markdown.name for r in results)
    per_industry = ["01-industry-brief", "04-personas", "06-playbook", "07-demo-storyline", "08-web-copy"]
    expected = [f"{a}-{i}.md" for a in per_industry for i in INDUSTRIES] + [
        "02-competitive-overview.md", "02-battlecard-comp-a.md", "03-strategy-check.md", "05-messaging-framework.md"]
    assert names == sorted(expected)
    for r in results:
        assert r.docx.is_file()
        Document(str(r.docx))  # opens without error
    assert (launch.dir / "outputs" / "05-messaging-grid.xlsx").is_file()
    assert (launch.dir / "outputs" / "images" / "positioning-map.png").is_file()
    for i in INDUSTRIES:
        assert (launch.dir / "outputs" / "images" / f"committee-{i}.png").is_file()


def test_every_asset_has_status_footer_and_verify_section(launch):
    for r in render(launch):
        md = r.markdown.read_text()
        assert "_Draft · Unbranded draft · Approved by PMM" in md
        assert FOOTER.format(date=load_yaml(launch.dir / "launch.yaml")["created_at"][:10]) in md
        assert "## Before you use this" in md and "## Sources" in md
        doc = Document(str(r.docx))
        footer = " ".join(p.text for p in doc.sections[0].footer.paragraphs)
        assert "Verify items marked [VERIFY] before external use." in footer
        header = " ".join(p.text for p in doc.sections[0].header.paragraphs)
        assert "Unbranded draft" in header


def test_markdown_tables_and_headings_become_word_structure(launch):
    render(launch)
    doc = Document(str(launch.dir / "outputs" / "01-industry-brief-ind-a.docx"))
    styles = [p.style.name for p in doc.paragraphs]
    assert styles[0] == "Title"
    assert "Heading 2" in styles and "List Bullet" in styles
    assert doc.tables, "the 'At a glance' table should be a real Word table"
    assert doc.tables[0].rows[0].cells[0].text == "Trend"
    assert len(doc.inline_shapes) == 0  # no image in the industry brief
    personas = Document(str(launch.dir / "outputs" / "04-personas-ind-a.docx"))
    assert len(personas.inline_shapes) == 1  # the committee diagram


def test_brand_fonts_and_colors_are_applied(launch):
    brand = f.brand()
    brand.update({"unbranded": False, "fonts": {"heading": "Georgia", "body": "Verdana",
                                                 "fallback_heading": "Arial", "fallback_body": "Arial"}})
    brand["colors"]["primary"] = "#123456"
    dump_yaml(brand, launch.dir / "brand.yaml")
    render(launch)
    doc = Document(str(launch.dir / "outputs" / "03-strategy-check.docx"))
    assert doc.styles["Normal"].font.name == "Verdana"
    assert doc.styles["Heading 1"].font.name == "Georgia"
    assert str(doc.styles["Heading 1"].font.color.rgb) == "123456"
    md = (launch.dir / "outputs" / "03-strategy-check.md").read_text()
    assert "Unbranded draft" not in md


def test_word_template_is_used_as_the_base(launch):
    make_word_template(launch.dir / "inputs" / "brand" / "brand.dotx", heading_font="Garamond")
    brand = f.brand()
    brand.update({"unbranded": False, "word_template": {
        "path": "inputs/brand/brand.dotx", "style_map": {"heading1": "Heading 1", "body": "Normal"}}})
    dump_yaml(brand, launch.dir / "brand.yaml")
    render(launch)
    doc = Document(str(launch.dir / "outputs" / "05-messaging-framework.docx"))
    assert doc.styles["Heading 1"].font.name == "Garamond"  # the template's style, untouched


def test_logo_goes_in_the_header(launch, tmp_path):
    import matplotlib

    matplotlib.use("Agg")
    import matplotlib.pyplot as plt

    fig = plt.figure(figsize=(1, 0.4))
    fig.savefig(launch.dir / "inputs" / "brand" / "logo.png")
    plt.close(fig)
    brand = f.brand()
    brand.update({"unbranded": False, "logo": {"light_bg": "inputs/brand/logo.png", "placement": "header-left"}})
    dump_yaml(brand, launch.dir / "brand.yaml")
    render(launch)
    doc = Document(str(launch.dir / "outputs" / "05-messaging-framework.docx"))
    header_xml = doc.sections[0].header._element.xml
    assert "graphicData" in header_xml


def test_auto_mode_assets_say_not_reviewed(tmp_path):
    s = LaunchState.create(tmp_path / "a", "a")
    s.data["industries"].append({"id": "ind-a", "name": "A"})
    s.set_auto_mode(True)
    s.refresh()
    for unit in ["intake", "research:ind-a"]:
        s.start(unit)
        write_outputs(s.dir, unit)
        s.submit(unit)
    s.save()
    md = Renderer(s.dir).render_all()[0].markdown.read_text()
    assert "Not reviewed by PMM" in md


def test_awaiting_approval_is_labeled(tmp_path):
    s = LaunchState.create(tmp_path / "w", "w")
    s.data["industries"].append({"id": "ind-a", "name": "A"})
    s.refresh()
    s.start("intake")
    write_outputs(s.dir, "intake")
    s.submit("intake")
    s.approve("intake")
    s.start("research:ind-a")
    write_outputs(s.dir, "research:ind-a")
    s.submit("research:ind-a")
    s.save()
    md = Renderer(s.dir).render_all()[0].markdown.read_text()
    assert "Awaiting PMM approval" in md


def test_verify_items_and_sources_are_listed(launch):
    doc = f.industry_research("ind-a")
    doc["trends"][0]["impact"] = "Costs may rise [VERIFY] check the regulator's estimate"
    doc["trends"][0]["so_what"] = "Placeholder [ASSUMPTION]"
    dump_yaml(doc, launch.dir / "research" / "industry-ind-a.yaml")
    render(launch)
    md = (launch.dir / "outputs" / "01-industry-brief-ind-a.md").read_text()
    assert "**[VERIFY]:** 1 · **[ASSUMPTION]:** 1" in md
    assert "- Costs may rise [VERIFY] check the regulator's estimate" in md
    assert "**[S1]** Source title, Publisher (2026-01-01" in md


def test_unregistered_citation_is_flagged(launch):
    doc = f.industry_research("ind-a")
    doc["trends"][0]["summary"] = "Claim [S99]"
    dump_yaml(doc, launch.dir / "research" / "industry-ind-a.yaml")
    render(launch)
    md = (launch.dir / "outputs" / "01-industry-brief-ind-a.md").read_text()
    assert "**[S99]** [VERIFY] not found in the source register" in md


def test_strategy_check_shows_agent_view_and_pmm_response(launch):
    doc = f.strategy_check()
    doc["points"][0]["pmm_response"] = {"decision": "adjust", "adjusted_direction": "Lead with the second buyer",
                                        "reason": "field feedback"}
    doc["resolved_direction"] = "Resolved placeholder direction"
    dump_yaml(doc, launch.dir / "strategy-check.yaml")
    render(launch)
    md = (launch.dir / "outputs" / "03-strategy-check.md").read_text()
    assert "**Agent's stance:** agree" in md
    assert "**PMM decision:** adjust" in md
    assert "**Adjusted direction:** Lead with the second buyer" in md
    assert "Resolved placeholder direction" in md


def test_rendering_is_deterministic(launch):
    first = {r.markdown.name: (r.markdown.read_text(), _docx_parts(r.docx)) for r in render(launch)}
    second = {r.markdown.name: (r.markdown.read_text(), _docx_parts(r.docx)) for r in render(launch)}
    assert first == second


def _docx_parts(path):
    with zipfile.ZipFile(path) as z:
        return {n: z.read(n) for n in z.namelist()}


def test_png_diagrams_are_deterministic(launch):
    render(launch)
    a = (launch.dir / "outputs" / "images" / "positioning-map.png").read_bytes()
    render(launch)
    assert (launch.dir / "outputs" / "images" / "positioning-map.png").read_bytes() == a


def test_label_and_citation_helpers():
    data = {"a": "x [VERIFY] y [S3]", "b": [{"source_id": "S10"}, "[PROOF NEEDED]"]}
    assert labels_in(data)["counts"] == {"verify": 1, "assumption": 0, "proof": 1}
    assert cited_in(data) == ["S3", "S10"]


def test_cli_render(launch, monkeypatch, capsys):
    from scripts import launch as cli

    monkeypatch.setattr(cli, "LAUNCHES_DIR", launch.dir.parent)
    assert cli.main(["render", launch.dir.name, "--unit", "messaging"]) == 0
    assert "05-messaging-framework.md" in capsys.readouterr().out


# ------------------------------------------------------------------ M4: demo, web, Excel grid
def test_excel_grid_has_filters_frozen_header_and_one_row_per_persona(launch):
    from openpyxl import load_workbook

    render(launch)
    wb = load_workbook(launch.dir / "outputs" / "05-messaging-grid.xlsx")
    ws = wb["Messaging grid"]
    assert ws.freeze_panes == "A2"
    assert ws.auto_filter.ref == "A1:J3"
    assert [c.value for c in ws[1]][:4] == ["Industry", "Persona", "Buying role", "Headline"]
    assert [ws.cell(r, 1).value for r in (2, 3)] == ["IND-A", "IND-B"]
    assert ws.cell(2, 9).value == "[ASSUMPTION]"
    assert ws.cell(1, 1).fill.fgColor.rgb.endswith("112233")  # brand primary
    assert "Verify items marked [VERIFY]" in ws.oddFooter.center.text
    assert "Draft generated with Industry Launch Accelerator" in wb["About"]["A3"].value


def test_excel_grid_uses_brand_font(launch):
    from openpyxl import load_workbook

    brand = f.brand()
    brand["fonts"]["body"] = "Verdana"
    dump_yaml(brand, launch.dir / "brand.yaml")
    render(launch)
    ws = load_workbook(launch.dir / "outputs" / "05-messaging-grid.xlsx")["Messaging grid"]
    assert {ws.cell(r, c).font.name for r in (1, 2) for c in (1, 4)} == {"Verdana"}


def test_demo_storyline_content(launch):
    doc = f.demo("ind-a")
    doc["moments"] = [{"id": "DM-1", "title": "Catch it early", "tell": "t", "show": "s", "tell_back": "b",
                       "pillar_id": "PL-1", "pain_ref": "P-1#1", "talk_track": "say this",
                       "verify": ["[VERIFY] confirm the product can flag this"]}]
    doc["aha"] = {"moment_id": "DM-1", "description": "the room sees the flag"}
    doc["cast"] = [{"persona_id": "P-1", "role_in_demo": "watches for effort"}]
    dump_yaml(doc, launch.dir / "demos" / "ind-a.yaml")
    render(launch)
    md = (launch.dir / "outputs" / "07-demo-storyline-ind-a.md").read_text()
    assert "### DM-1: Catch it early  (aha)" in md
    assert "| DM-1 | Catch it early | PL-1 | r: placeholder pain |" in md
    assert "## 5-minute version" in md and "## 20-minute version" in md
    assert "[VERIFY] confirm the product can flag this" in md


def test_web_copy_content(launch):
    doc = f.web("ind-a")
    doc["sections"] = [{"id": "WS-1", "name": "Why now", "purpose": "urgency", "copy": "Draft copy here.",
                        "messaging_refs": ["PL-1", "P-1"]}]
    doc["headline_variants"] = ["Variant one", "Variant two"]
    dump_yaml(doc, launch.dir / "web" / "ind-a.yaml")
    render(launch)
    md = (launch.dir / "outputs" / "08-web-copy-ind-a.md").read_text()
    assert "| WS-1 | Why now | urgency | PL-1, P-1 |" in md
    assert "### WS-1: Why now\n\nDraft copy here." in md  # the section text itself, not a Python object
    doc = Document(str(launch.dir / "outputs" / "08-web-copy-ind-a.docx"))
    assert "Draft copy here." in [p.text for p in doc.paragraphs]
    assert "1. Variant one" in md and "2. Variant two" in md
    assert "- t [VERIFY]" in md and "Search volume has not been checked" in md


# ------------------------------------------------------------------ regression: data fields named like built-ins
def test_fields_named_like_dict_methods_render_as_data():
    from scripts.render.render import DataFirstEnvironment

    env = DataFirstEnvironment()
    t = env.from_string("{{ d.copy }}|{{ d.items }}|{{ d.values }}|{{ d.count }}|{{ d.missing }}")
    assert t.render(d={"copy": "C", "items": "I", "values": "V", "count": 3}) == "C|I|V|3|"


def test_no_python_internals_in_any_rendered_file(launch):
    doc = f.web("ind-a")
    doc["sections"] = [{"id": "WS-1", "name": "Why now", "purpose": "p", "copy": "Section text.", "messaging_refs": []}]
    dump_yaml(doc, launch.dir / "web" / "ind-a.yaml")
    for r in render(launch):
        text = r.markdown.read_text()
        for artifact in ("<built-in method", "<bound method", "<function ", " object at 0x"):
            assert artifact not in text, f"{r.markdown.name}: {artifact}"
        body = " ".join(p.text for p in Document(str(r.docx)).paragraphs)
        assert "<built-in method" not in body and " object at 0x" not in body, r.docx.name


def test_word_files_use_the_current_format_not_compatibility_mode(launch):
    render(launch)
    with zipfile.ZipFile(launch.dir / "outputs" / "03-strategy-check.docx") as z:
        settings = z.read("word/settings.xml").decode()
    assert 'w:name="compatibilityMode"' in settings
    assert 'w:name="compatibilityMode" w:uri="http://schemas.microsoft.com/office/word" w:val="15"' in settings


def test_committee_diagram_handles_large_committees_and_repeated_steps(tmp_path):
    """Regression: 6+ people used to shrink the text and push arrows off the image."""
    from scripts.brand.theme import NEUTRAL_COLORS, NEUTRAL_FONTS
    from scripts.render.diagrams import committee_map

    ids = [f"P-{n}" for n in range(1, 7)]
    committee = {"industry": "x", "decision_sequence": ["P-1", "P-2", "P-3", "P-4", "P-5", "P-2"],
                 "edges": [{"from": "P-1", "to": "P-2", "influence": "high"},
                           {"from": "P-6", "to": "P-1", "influence": "medium"},
                           {"from": "P-4", "to": "P-5", "influence": "low"}]}
    roles = {pid: f"A fairly long job title number {pid}" for pid in ids}
    brand = {"colors": NEUTRAL_COLORS, "fonts": NEUTRAL_FONTS}
    path = committee_map(committee, roles, brand, tmp_path / "c.png")
    from matplotlib import image

    h, w = image.imread(path).shape[:2]
    assert w > 1000 and h > 500  # big enough to stay readable on a page
