"""Build tiny .docx/.pptx/.pdf files inside tests, so no binary samples live in the repo."""

from __future__ import annotations

from pathlib import Path


def make_docx(path: Path, heading: str, paragraph: str, table: list[list[str]] | None = None) -> Path:
    from docx import Document

    doc = Document()
    doc.add_heading(heading, level=1)
    doc.add_paragraph(paragraph)
    if table:
        t = doc.add_table(rows=len(table), cols=len(table[0]))
        for r, row in enumerate(table):
            for c, value in enumerate(row):
                t.cell(r, c).text = value
    doc.save(str(path))
    return path


def make_pptx(path: Path, title: str, notes: str = "") -> Path:
    from pptx import Presentation

    prs = Presentation()
    slide = prs.slides.add_slide(prs.slide_layouts[1])
    slide.shapes.title.text = title
    if notes:
        slide.notes_slide.notes_text_frame.text = notes
    prs.save(str(path))
    return path


def make_pdf(path: Path, text: str) -> Path:
    """A one-page PDF containing `text` (ASCII only), written by hand."""
    stream = f"BT /F1 12 Tf 72 720 Td ({text}) Tj ET".encode()
    objects = [
        b"<< /Type /Catalog /Pages 2 0 R >>",
        b"<< /Type /Pages /Kids [3 0 R] /Count 1 >>",
        b"<< /Type /Page /Parent 2 0 R /MediaBox [0 0 612 792] /Contents 4 0 R "
        b"/Resources << /Font << /F1 5 0 R >> >> >>",
        b"<< /Length %d >>\nstream\n" % len(stream) + stream + b"\nendstream",
        b"<< /Type /Font /Subtype /Type1 /BaseFont /Helvetica >>",
    ]
    out = bytearray(b"%PDF-1.4\n")
    offsets = []
    for i, obj in enumerate(objects, start=1):
        offsets.append(len(out))
        out += b"%d 0 obj\n" % i + obj + b"\nendobj\n"
    xref = len(out)
    out += b"xref\n0 %d\n0000000000 65535 f \n" % (len(objects) + 1)
    for off in offsets:
        out += b"%010d 00000 n \n" % off
    out += b"trailer\n<< /Size %d /Root 1 0 R >>\nstartxref\n%d\n%%%%EOF\n" % (len(objects) + 1, xref)
    path.write_bytes(bytes(out))
    return path


def make_word_template(path: Path, heading_font: str = "Georgia", heading_color: str = "123456") -> Path:
    """A .docx (or .dotx, by extension) whose Heading 1 style uses a given font and color."""
    import io
    import zipfile

    from docx import Document
    from docx.shared import RGBColor

    doc = Document()
    style = doc.styles["Heading 1"]
    style.font.name = heading_font
    style.font.color.rgb = RGBColor.from_string(heading_color)
    if path.suffix.lower() != ".dotx":
        doc.save(str(path))
        return path
    buffer = io.BytesIO()
    doc.save(buffer)
    with zipfile.ZipFile(io.BytesIO(buffer.getvalue())) as src, zipfile.ZipFile(path, "w") as dst:
        for item in src.infolist():
            data = src.read(item.filename)
            if item.filename == "[Content_Types].xml":
                data = data.replace(
                    b"wordprocessingml.document.main+xml", b"wordprocessingml.template.main+xml"
                )
            dst.writestr(item, data)
    return path
